"""
Multi-sheet xlsx from structured JSON (openpyxl).

Input JSON (stdin or ``--input`` file):
  title: optional str
  notes: optional str — cover sheet body
  sheets: required list of {
    name: str,
    columns: list[str],
    rows: list[list] — row-major cell values,
    formats: optional list[str|none] — per-column number formats (Excel codes), length may match columns
  }
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _load_payload(source: Path | None) -> dict[str, Any]:
    raw = source.read_text(encoding="utf-8") if source else sys.stdin.read()
    data = json.loads(raw)
    if "sheets" not in data or not isinstance(data["sheets"], list):
        raise ValueError("payload must include 'sheets' array")
    return data


def render_xlsx(payload: dict[str, Any], output_path: Path) -> None:
    title = payload.get("title") or ""
    notes = payload.get("notes") or ""
    sheets_spec: list[dict[str, Any]] = payload["sheets"]

    if not sheets_spec and not (title or notes):
        raise ValueError("payload must include at least one sheet or cover content")

    wb = Workbook()
    used_names: set[str] = set()

    def unique_sheet_name(requested: str) -> str:
        base = (requested or "Sheet").strip() or "Sheet"
        base = base[:31]
        name = base
        n = 2
        while name.lower() in used_names:
            suffix = f" ({n})"
            name = (base[: 31 - len(suffix)] + suffix).strip()
            n += 1
        used_names.add(name.lower())
        return name

    if title or notes:
        cover = wb.active
        cover.title = unique_sheet_name("Cover")
        cover["A1"] = title or "Report"
        cover["A1"].font = Font(bold=True, size=14)
        if notes:
            cover["A3"] = notes

    for i, spec in enumerate(sheets_spec):
        name = unique_sheet_name(str(spec.get("name") or "Sheet"))
        if i == 0 and not (title or notes):
            ws = wb.active
            ws.title = name
        else:
            ws = wb.create_sheet(name)

        columns = spec.get("columns") or []
        formats = spec.get("formats")

        for col_idx, heading in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=heading)
            cell.font = Font(bold=True)

        rows = spec.get("rows") or []
        for r_idx, row in enumerate(rows, start=2):
            if not isinstance(row, (list, tuple)):
                row = [row]
            for c_idx, val in enumerate(row, start=1):
                if c_idx > len(columns):
                    break
                ws.cell(row=r_idx, column=c_idx, value=val)

        if formats and isinstance(formats, list):
            for c_idx, _ in enumerate(columns, start=1):
                if c_idx - 1 >= len(formats):
                    break
                code = formats[c_idx - 1]
                if not code:
                    continue
                letter = get_column_letter(c_idx)
                for row_i in range(2, ws.max_row + 1):
                    ws[f"{letter}{row_i}"].number_format = str(code)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Render spreadsheet JSON to xlsx.")
    p.add_argument("--input", "-i", type=Path, help="Input JSON file (default: stdin)")
    p.add_argument("--output", "-o", type=Path, required=True, help="Output .xlsx path")
    args = p.parse_args(argv)

    payload = _load_payload(args.input)
    render_xlsx(payload, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
