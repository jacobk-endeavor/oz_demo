"""Smoke tests for sandboxTemplates (run from repo: ``PYTHONPATH=backend/oz python -m unittest discover -s backend/oz/sandboxTemplates``)."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

OZ_ROOT = Path(__file__).resolve().parents[1]
if str(OZ_ROOT) not in sys.path:
    sys.path.insert(0, str(OZ_ROOT))

from sandboxTemplates.report_docx import render_docx  # noqa: E402
from sandboxTemplates.report_pdf import render_pdf  # noqa: E402
from sandboxTemplates.spreadsheet import render_xlsx  # noqa: E402


class TestSpreadsheet(unittest.TestCase):
    def test_multi_sheet_and_cover(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "o.xlsx"
            render_xlsx(
                {
                    "title": "T",
                    "notes": "N",
                    "sheets": [
                        {
                            "name": "Data",
                            "columns": ["SKU", "Qty"],
                            "rows": [["A1", 2]],
                            "formats": [None, "0"],
                        }
                    ],
                },
                out,
            )
            self.assertTrue(out.is_file())
            wb = load_workbook(out)
            names = wb.sheetnames
            self.assertEqual(names[0], "Cover")
            self.assertIn("Data", names)

    def test_no_cover_reuses_active(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "o.xlsx"
            render_xlsx(
                {
                    "sheets": [
                        {"name": "Only", "columns": ["x"], "rows": [[1]]},
                    ],
                },
                out,
            )
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Only"])


class TestReports(unittest.TestCase):
    def test_docx_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "r.docx"
            render_docx(
                {
                    "title": "Hello",
                    "sections": [
                        {"type": "paragraph", "text": "Body line."},
                        {"type": "table", "headers": ["a", "b"], "rows": [[1, 2]]},
                    ],
                },
                out,
            )
            self.assertGreater(out.stat().st_size, 2000)
            from docx import Document as Doc

            d = Doc(str(out))
            texts = [p.text for p in d.paragraphs if p.text.strip()]
            joined = " ".join(texts)
            self.assertIn("Hello", joined)
            self.assertIn("Body line.", joined)

    def test_pdf_writes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "r.pdf"
            render_pdf(
                {
                    "title": "Hi",
                    "sections": [{"type": "paragraph", "text": "PDF body."}],
                },
                out,
            )
            self.assertGreater(out.stat().st_size, 500)
            self.assertEqual(out.read_bytes()[:4], b"%PDF")


class TestModuleCli(unittest.TestCase):
    def _run_mod(self, mod: str, payload: dict, suffix: str) -> Path:
        td = tempfile.mkdtemp()
        self.addCleanup(lambda: shutil.rmtree(td, ignore_errors=True))
        inp = Path(td) / "in.json"
        outp = Path(td) / f"out{suffix}"
        inp.write_text(json.dumps(payload), encoding="utf-8")
        env = {**os.environ, "PYTHONPATH": str(OZ_ROOT)}
        r = subprocess.run(
            [sys.executable, "-m", f"sandboxTemplates.{mod}", "--input", str(inp), "--output", str(outp)],
            cwd=str(OZ_ROOT),
            env=env,
            capture_output=True,
            text=True,
            timeout=60,
        )
        self.assertEqual(r.returncode, 0, msg=r.stderr + r.stdout)
        self.assertTrue(outp.is_file())
        return outp

    def test_cli_spreadsheet(self) -> None:
        path = self._run_mod(
            "spreadsheet",
            {"sheets": [{"name": "X", "columns": ["c"], "rows": [["v"]]}]},
            ".xlsx",
        )
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["X"])

    def test_cli_docx(self) -> None:
        path = self._run_mod(
            "report_docx",
            {"title": "T", "sections": [{"type": "paragraph", "text": "x"}]},
            ".docx",
        )
        self.assertGreater(path.stat().st_size, 1000)

    def test_cli_pdf(self) -> None:
        path = self._run_mod(
            "report_pdf",
            {"title": "T", "sections": [{"type": "paragraph", "text": "y"}]},
            ".pdf",
        )
        self.assertEqual(path.read_bytes()[:4], b"%PDF")


if __name__ == "__main__":
    unittest.main()
