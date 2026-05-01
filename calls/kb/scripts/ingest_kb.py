#!/usr/bin/env python3
"""
Ingest raw files (PDF, Excel, PPTX, text, JSON) into kb_sources + kb_rag_chunks.

Loads repo-root `.env` (override=False) like calls/sauron/scripts/ingest_calls_pgvector.py.

Usage:
  cd calls/kb/scripts
  python3 -m venv .venv && .venv/bin/pip install -r requirements-kb-ingest.txt
  .venv/bin/python ingest_kb.py <path-or-dir> [options]
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Literal
from urllib.parse import quote_plus

SCRIPT_DIR = Path(__file__).resolve().parent
CALLS_DIR = SCRIPT_DIR.parent.parent
REPO_ROOT = CALLS_DIR.parent
SCHEMA_PATH = SCRIPT_DIR / "schema.sql"
KB_SCHEMAS_DIR = REPO_ROOT / "kb_schemas"

CHUNK_CHARS = 1400
CHUNK_OVERLAP = 200
EMBED_MODEL = "text-embedding-3-small"
EMBED_DIM = 1536
EMBED_PRICE_USD_PER_1K = 0.00002
PDF_MIN_AVG_CHARS = 50
DEFAULT_MAX_EMBED_SPEND_USD = 10.0
EXCEL_ROW_CHUNK = 20

AssetKind = Literal["pdf", "excel", "pptx", "text", "image", "structured"]
TypeFilter = Literal["auto", "pdf", "excel", "text", "image"]


def load_dotenv(path: Path) -> None:
    if not path.is_file():
        return
    try:
        from dotenv import load_dotenv as _load

        _load(path, override=False)
    except ImportError:
        for raw in path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, val = line.partition("=")
            key, val = key.strip(), val.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = val


def build_conninfo() -> str:
    url = (os.environ.get("DATABASE_URL") or "").strip()
    if url:
        return url
    host = os.environ.get("PGHOST", "").strip()
    user = os.environ.get("PGUSER", "").strip()
    password = os.environ.get("PGPASSWORD", "").strip()
    db = (os.environ.get("PGDATABASE") or "defaultdb").strip()
    port = os.environ.get("PGPORT", "5432").strip()
    sslmode = os.environ.get("PGSSLMODE", "require").strip()
    if not all([host, user, password, db]):
        raise SystemExit(
            "Missing DB config: set DATABASE_URL or PGHOST, PGUSER, PGPASSWORD, PGDATABASE "
            "(optional PGPORT, PGSSLMODE)."
        )
    pw = quote_plus(password)
    return f"postgresql://{user}:{pw}@{host}:{port}/{db}?sslmode={sslmode}"


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def source_id_from_sha256(hex_digest: str) -> str:
    return hex_digest[:12].lower()


def chunk_transcript(text: str, size: int, overlap: int) -> list[str]:
    text = (text or "").strip()
    if not text:
        return []
    if overlap >= size:
        overlap = max(0, size // 5)
    out: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        out.append(text[start:end])
        if end >= n:
            break
        start = end - overlap
    return out


def guess_mime(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    if mime:
        return mime
    ext = path.suffix.lower()
    if ext == ".pdf":
        return "application/pdf"
    if ext in (".xlsx", ".xlsm"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if ext == ".xls":
        return "application/vnd.ms-excel"
    if ext == ".pptx":
        return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    if ext in (".txt", ".md"):
        return "text/plain"
    if ext == ".json":
        return "application/json"
    return "application/octet-stream"


def classify_asset_kind(path: Path, mime: str, type_arg: TypeFilter) -> AssetKind | None:
    name = path.name.lower()
    m = mime.lower()
    ext = path.suffix.lower()
    if type_arg == "pdf":
        return "pdf" if name.endswith(".pdf") or "pdf" in m else None
    if type_arg == "excel":
        return (
            "excel"
            if name.endswith((".xlsx", ".xls", ".xlsm", ".xlsb", ".ods", ".csv"))
            or "spreadsheet" in m
            or "excel" in m
            else None
        )
    if type_arg == "text":
        if ext in (".txt", ".md", ".markdown", ".csv", ".json") or m.startswith("text/") or m == "application/json":
            return "text"
        return None
    if type_arg == "image":
        if m.startswith("image/") or ext in ("png", "jpg", "jpeg", "gif", "webp", "svg", "bmp", "tif", "tiff", "heic"):
            return "image"
        return None
    # auto
    if name.endswith(".pptx") or "presentationml" in m:
        return "pptx"
    if name.endswith(".pdf") or m == "application/pdf":
        return "pdf"
    if name.endswith((".xlsx", ".xls", ".xlsm", ".xlsb", ".ods")) or "spreadsheet" in m or "excel" in m:
        return "excel"
    if m.startswith("image/") or name.split(".")[-1] in ("png", "jpg", "jpeg", "gif", "webp", "svg", "bmp", "tif", "tiff"):
        return "image"
    if name.endswith(".json") and re.search(r"product_catalog|recommendations", name, re.I):
        return "structured"
    if name.endswith((".txt", ".md", ".markdown", ".csv", ".json")) or m.startswith("text/") or m == "application/json":
        return "text"
    return "text"


def slug_part(value: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9_-]+", "-", value.strip())
    s = re.sub(r"-+", "-", s).strip("-")
    return s or "unknown"


def slugify_sheet_name(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-")
    return s or "sheet"


def csv_escape_cell(cell: str) -> str:
    if not re.search(r'[",\n]', cell):
        return cell
    return '"' + cell.replace('"', '""') + '"'


def row_to_csv(row: list[str]) -> str:
    return ",".join(csv_escape_cell(c) for c in row)


def pad_row(row: list[str], width: int) -> list[str]:
    out = list(row)
    while len(out) < width:
        out.append("")
    return out


FULL_TEXT_DOC_KINDS = frozenset(
    {"marketing", "install", "tech-bulletin", "warranty", "order-guide", "presentation"}
)


@dataclass
class DocKindResult:
    doc_kind: str
    confidence: float
    matched_by: str
    signals: list[str]
    brand: str | None
    product_line: str | None
    year: int | None
    distributor_branded: bool


FILENAME_RULES: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"\btechnical bulletin\b", re.I), "tech-bulletin"),
    (re.compile(r"\bmaster spec\b", re.I), "master-spec"),
    (re.compile(r"\border guide\b", re.I), "order-guide"),
    (re.compile(r"\bsku list\b", re.I), "order-guide"),
    (re.compile(r"\bwarranty\b", re.I), "warranty"),
    (re.compile(r"\bline card\b", re.I), "catalog"),
    (re.compile(r"\bproduct catalog\b", re.I), "catalog"),
    (re.compile(r"\bproduct[- ]offerings\b", re.I), "catalog"),
    (re.compile(r"\bsell sheet\b", re.I), "spec-sheet"),
    (re.compile(r"\bsales sheet\b", re.I), "spec-sheet"),
    (re.compile(r"\binstall guide\b", re.I), "install"),
    (re.compile(r"\binstallation instructions\b", re.I), "install"),
    (re.compile(r"\binstall instructions\b", re.I), "install"),
    (re.compile(r"\bcolor comparison\b", re.I), "visual-catalog"),
    (re.compile(r"\bcolor chart\b", re.I), "visual-catalog"),
    (re.compile(r"\bbrochure\b", re.I), "marketing"),
    (re.compile(r"\boverview\b", re.I), "marketing"),
    (re.compile(r"\bpresentation\b", re.I), "presentation"),
]

CONTENT_RULES: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"\binstallation instructions\b", re.I), "install"),
    (re.compile(r"\binstall guide\b", re.I), "install"),
    (re.compile(r"\bwarranty\b", re.I), "warranty"),
    (re.compile(r"\bproduct catalog\b", re.I), "catalog"),
    (re.compile(r"\bline card\b", re.I), "catalog"),
    (re.compile(r"\btechnical bulletin\b", re.I), "tech-bulletin"),
    (re.compile(r"\bsell sheet\b", re.I), "spec-sheet"),
    (re.compile(r"\bsales sheet\b", re.I), "spec-sheet"),
    (re.compile(r"\bmaster spec\b", re.I), "master-spec"),
    (re.compile(r"\border guide\b", re.I), "order-guide"),
]

BRANDS = ["AZEK", "Deckorators", "TimberTech", "Trex", "Fiberon", "Millboard", "Russin", "TFP"]
PRODUCT_LINES = ["Captivate", "Evolution", "Black Label", "Maximo Thermo", "Shadow Line+"]


def classify_doc_kind(path: Path, mime: str, first_page_text: str = "") -> DocKindResult:
    basename = path.name
    normalized_basename = re.sub(r"[_+.]+", " ", basename)
    ext = path.suffix.lower()
    m = mime.lower()
    context = f"{basename}\n{first_page_text}"

    if ext == ".pptx" or "presentation" in m:
        return DocKindResult(
            "presentation",
            1.0,
            "filename",
            ["ext:pptx" if ext == ".pptx" else "mime:presentation"],
            _pick_brand(context),
            _pick_product_line(context),
            _pick_year(context),
            bool(re.search(r"russin logo|with russin", context, re.I)),
        )

    excel_like = ext in (".xlsx", ".xls", ".xlsm", ".xlsb", ".ods", ".csv")
    if excel_like or "spreadsheet" in m or "excel" in m:
        return DocKindResult(
            "tabular-reference",
            1.0,
            "filename",
            [f"ext:{ext or 'unknown'}"],
            _pick_brand(context),
            _pick_product_line(context),
            _pick_year(context),
            bool(re.search(r"russin logo|with russin", context, re.I)),
        )

    if ext == ".json" and re.search(r"product_catalog|recommendations", basename, re.I):
        return DocKindResult(
            "structured-data",
            1.0,
            "filename",
            ["filename:known-json-structured-source"],
            _pick_brand(context),
            _pick_product_line(context),
            _pick_year(context),
            bool(re.search(r"russin logo|with russin", context, re.I)),
        )

    for rx, dk in FILENAME_RULES:
        if rx.search(normalized_basename):
            return DocKindResult(
                dk,
                1.0,
                "filename",
                [rx.pattern],
                _pick_brand(context),
                _pick_product_line(context),
                _pick_year(context),
                bool(re.search(r"russin logo|with russin", context, re.I)),
            )

    fp = first_page_text.strip()
    if fp:
        for rx, dk in CONTENT_RULES:
            if rx.search(fp):
                return DocKindResult(
                    dk,
                    0.8,
                    "content",
                    [rx.pattern],
                    _pick_brand(context),
                    _pick_product_line(context),
                    _pick_year(context),
                    bool(re.search(r"russin logo|with russin", context, re.I)),
                )

    return DocKindResult(
        "unknown",
        0.0,
        "fallback",
        ["no-match"],
        _pick_brand(context),
        _pick_product_line(context),
        _pick_year(context),
        bool(re.search(r"russin logo|with russin", context, re.I)),
    )


def _pick_brand(text: str) -> str | None:
    low = text.lower()
    for b in BRANDS:
        if b.lower() in low:
            return b
    return None


def _pick_product_line(text: str) -> str | None:
    low = text.lower()
    for pl in PRODUCT_LINES:
        if pl.lower() in low:
            return pl
    return None


def _pick_year(text: str) -> int | None:
    m = re.search(r"\b(20\d{2})\b", text)
    if not m:
        return None
    y = int(m.group(1))
    if 2023 <= y <= 2026:
        return y
    return None


def extract_pdf_pages(data: bytes, password: str | None) -> tuple[list[str], Literal["ok", "needs_ocr", "encrypted"]]:
    """Extract PDF pages with pypdf then pdfminer fallback (same thresholds as frontend/pdfExtractor.ts)."""
    from pdfminer.high_level import extract_text
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError

    pw = password or None
    try:
        reader = PdfReader(io.BytesIO(data), password=pw)
    except FileNotDecryptedError:
        return [], "encrypted"
    except Exception as e:
        if "password" in str(e).lower() or "encrypt" in str(e).lower():
            return [], "encrypted"
        raise

    def avg_chars(pages: list[str]) -> float:
        if not pages:
            return 0.0
        return sum(len(p) for p in pages) / len(pages)

    primary: list[str] = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        primary.append(t.replace("\r\n", "\n").strip())

    if len(primary) > 0 and avg_chars(primary) >= PDF_MIN_AVG_CHARS:
        return primary, "ok"

    n = len(primary)
    fallback: list[str] = []
    for i in range(n):
        try:
            if pw:
                t = extract_text(io.BytesIO(data), page_numbers=[i], password=pw) or ""
            else:
                t = extract_text(io.BytesIO(data), page_numbers=[i]) or ""
        except Exception:
            t = ""
        fallback.append(t.replace("\r\n", "\n").strip())

    if len(fallback) > 0 and avg_chars(fallback) >= PDF_MIN_AVG_CHARS:
        return fallback, "ok"

    return (fallback if fallback else primary), "needs_ocr"


def title_from_path(path: Path) -> str:
    return path.stem


def build_pdf_chunk_rows(
    *,
    source_id: str,
    title: str,
    pages: list[str],
    scope: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    chunk_index = 0
    for page_num, page_text in enumerate(pages, start=1):
        locator = f"page={page_num}"
        parts = chunk_transcript(page_text, CHUNK_CHARS, CHUNK_OVERLAP)
        for pi, part in enumerate(parts):
            header = f"[source={title}][{locator}]\n\n"
            content = header + part
            cid = f"{source_id}_p{page_num:03d}_{pi:05d}"
            rows.append(
                {
                    "chunk_id": cid,
                    "source_id": source_id,
                    "scope": scope,
                    "locator": locator,
                    "chunk_index": chunk_index,
                    "content": content,
                    "meta": {"kind": "pdf", "page": page_num, "part": pi},
                }
            )
            chunk_index += 1
    return rows


def extract_excel_units(source_id: str, data: bytes) -> tuple[list[dict[str, Any]], int, list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        units: list[dict[str, Any]] = []
        failed_sheets: list[str] = []
        sheet_names = list(wb.sheetnames)
        for sheet_name in sheet_names:
            try:
                ws = wb[sheet_name]
                rows_iter = ws.iter_rows(values_only=True)
                table_rows: list[list[str]] = []
                for row in rows_iter:
                    table_rows.append(["" if c is None else str(c) for c in row])
                if not table_rows:
                    continue
                headers = table_rows[0]
                body = table_rows[1:]
                width = max(len(headers), max((len(r) for r in body), default=0))
                headers = pad_row(list(headers), width)
                sheet_slug = slugify_sheet_name(sheet_name)
                for row_index in range(0, len(body), EXCEL_ROW_CHUNK):
                    start = row_index + 1
                    end = min(row_index + EXCEL_ROW_CHUNK, len(body))
                    slice_rows = body[row_index:end]
                    body_lines = [row_to_csv(pad_row(list(r), width)) for r in slice_rows]
                    header_line = row_to_csv(headers)
                    body_csv = "\n".join([header_line, *body_lines])
                    range_suffix = f"r{start:05d}_{end:05d}"
                    locator = f"sheet={sheet_name} rows={start}-{end}"
                    cid = f"{source_id}_{sheet_slug}_{range_suffix}"
                    units.append(
                        {
                            "chunk_id": cid,
                            "locator": locator,
                            "body": body_csv,
                            "meta": {"kind": "excel", "sheet": sheet_name},
                        }
                    )
            except Exception:
                failed_sheets.append(sheet_name)
        return units, len(sheet_names), failed_sheets
    finally:
        wb.close()


def extract_pptx_slides(data: bytes) -> list[tuple[int, str]]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides_out: list[tuple[int, str]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        title_shape = getattr(slide.shapes, "title", None)
        if title_shape is not None and getattr(title_shape, "text", None):
            texts.append(title_shape.text.strip())
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape is not title_shape:
                t = shape.text.strip()
                if t:
                    texts.append(t)
        if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
            nt = slide.notes_slide.notes_text_frame.text.strip()
            if nt:
                texts.append(nt)
        body = "\n\n".join(texts)
        slides_out.append((idx, body))
    return slides_out


def build_pptx_chunk_rows(
    *,
    source_id: str,
    title: str,
    slides: list[tuple[int, str]],
    scope: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    chunk_index = 0
    for slide_num, slide_text in slides:
        locator = f"slide={slide_num}"
        parts = chunk_transcript(slide_text, CHUNK_CHARS, CHUNK_OVERLAP)
        for pi, part in enumerate(parts):
            header = f"[source={title}][{locator}]\n\n"
            content = header + part
            cid = f"{source_id}_s{slide_num:03d}_{pi:05d}"
            rows.append(
                {
                    "chunk_id": cid,
                    "source_id": source_id,
                    "scope": scope,
                    "locator": locator,
                    "chunk_index": chunk_index,
                    "content": content,
                    "meta": {"kind": "pptx", "slide": slide_num, "part": pi},
                }
            )
            chunk_index += 1
    return rows


def validate_json_schema(instance: Any, schema_path: Path) -> None:
    import jsonschema

    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    jsonschema.validate(instance=instance, schema=schema)


def structured_units_from_json(file_path: Path, body: str) -> tuple[list[dict[str, Any]], str]:
    stem = file_path.stem.lower()
    data = json.loads(body)
    if stem == "product_catalog":
        validate_json_schema(data, KB_SCHEMAS_DIR / "product_catalog.schema.json")
        units = _product_catalog_units(data)
        return units, "product_catalog"
    if stem == "recommendations":
        validate_json_schema(data, KB_SCHEMAS_DIR / "recommendations.schema.json")
        units = _recommendations_units(data)
        return units, "recommendations"
    raise ValueError("no_schema")


def _as_str(v: Any, fb: str = "") -> str:
    return v.strip() if isinstance(v, str) and v.strip() else fb


def _as_num(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return float(v)
    if isinstance(v, float):
        return v
    return None


def _product_catalog_units(payload: dict[str, Any]) -> list[dict[str, Any]]:
    items_raw = payload.get("items")
    if not isinstance(items_raw, list):
        return []
    items = [x for x in items_raw if isinstance(x, dict)]
    units: list[dict[str, Any]] = []
    by_sub: dict[str, list[dict[str, Any]]] = {}
    by_line: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        sku = _as_str(item.get("sku"), "unknown")
        desc = _as_str(item.get("description"), "Unknown product")
        line = _as_str(item.get("product_line_code") or item.get("product_line"), "unknown")
        sub = _as_str(item.get("sub_category"), "unknown")
        uom = _as_str(item.get("uom"), "unknown")
        unit_price = _as_num(item.get("unit_price_median")) or _as_num(item.get("unit_price_avg"))
        qty = _as_num(item.get("total_qty_sold"))
        sales = _as_num(item.get("total_sales"))
        locator = f"[catalog:sku={sku}]"
        chunk_id = f"cat_sku_{slug_part(sku)}"
        body = (
            f"{locator}\n{desc} - sub-category {sub} in product line {line}. "
            f"UoM {uom}. unit_price ~{unit_price or 'n/a'}, qty_sold {qty or 'n/a'}, total_sales {sales or 'n/a'}."
        )
        units.append({"chunk_id": chunk_id, "locator": locator, "body": body})
        by_sub.setdefault(sub, []).append(item)
        by_line.setdefault(line, []).append(item)

    for sub, group in sorted(by_sub.items(), key=lambda x: x[0]):
        sku_count = len(group)
        qty_total = sum(_as_num(r.get("total_qty_sold")) or 0 for r in group)
        sales_total = sum(_as_num(r.get("total_sales")) or 0 for r in group)
        locator = f"[catalog:sub_category={sub}]"
        chunk_id = f"cat_subcat_{slug_part(sub)}"
        body = f"{locator}\nSub-category {sub}: {sku_count} SKUs, total_qty_sold {qty_total}, total_sales {sales_total}."
        units.append({"chunk_id": chunk_id, "locator": locator, "body": body})

    for line, group in sorted(by_line.items(), key=lambda x: x[0]):
        sku_count = len(group)
        qty_total = sum(_as_num(r.get("total_qty_sold")) or 0 for r in group)
        sales_total = sum(_as_num(r.get("total_sales")) or 0 for r in group)
        locator = f"[catalog:line={line}]"
        chunk_id = f"cat_line_{slug_part(line)}"
        body = f"{locator}\nProduct line {line}: {sku_count} SKUs, total_qty_sold {qty_total}, total_sales {sales_total}."
        units.append({"chunk_id": chunk_id, "locator": locator, "body": body})

    return units


def _recommendations_units(payload: dict[str, Any]) -> list[dict[str, Any]]:
    groups = [
        ("cross_sell", "xs"),
        ("upsell", "up"),
        ("margin_substitution", "ms"),
    ]
    units: list[dict[str, Any]] = []
    for key, short in groups:
        rules_raw = payload.get(key)
        if not isinstance(rules_raw, list):
            continue
        for index, rule in enumerate(rules_raw):
            if not isinstance(rule, dict):
                continue
            left = _as_str(rule.get("left_sku"), "unknown")
            right = _as_str(rule.get("right_sku"), "unknown")
            co_inv = _as_num(rule.get("co_invoices"))
            conf = _as_num(rule.get("confidence"))
            lift = _as_num(rule.get("lift"))
            locator = f"[recs:{key}:{left}#{index}]"
            chunk_id = f"recs_{short}_{slug_part(left)}_{index}"
            body = (
                f"{locator}\n{key.replace('_', ' ')} rule: customers who buy {left} also buy {right}. "
                f"co_invoices={co_inv or 'n/a'}, confidence={conf or 'n/a'}, lift={lift or 'n/a'}."
            )
            units.append({"chunk_id": chunk_id, "locator": locator, "body": body})
    return units


def embed_batches(client: Any, model: str, texts: list[str], batch_size: int) -> tuple[list[list[float]], int]:
    """Returns (vectors, total_tokens_estimated)."""
    all_vecs: list[list[float]] = []
    token_estimate = 0
    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        resp = client.embeddings.create(model=model, input=batch)
        ordered = sorted(resp.data, key=lambda x: x.index)
        for item in ordered:
            all_vecs.append(item.embedding)
        token_estimate += sum(max(1, len(t) // 4) for t in batch)
    return all_vecs, token_estimate


def delete_chunks_for_source(cur: Any, source_id: str) -> None:
    cur.execute("DELETE FROM kb_rag_chunks WHERE source_id = %s", (source_id,))


def fetch_source_row(cur: Any, source_id: str) -> dict[str, Any] | None:
    cur.execute(
        "SELECT source_id, path, first_seen_path, sha256, status, meta FROM kb_sources WHERE source_id = %s",
        (source_id,),
    )
    row = cur.fetchone()
    if row is None:
        return None
    return {
        "source_id": row[0],
        "path": row[1],
        "first_seen_path": row[2],
        "sha256": row[3],
        "status": row[4],
        "meta": row[5],
    }


def upsert_source_row(
    cur: Any,
    *,
    source_id: str,
    path: str,
    first_seen_path: str,
    sha256: str,
    mime: str,
    bytes_: int,
    page_count: int | None,
    sheet_count: int | None,
    status: str,
    meta: dict[str, Any],
) -> None:
    cur.execute(
        """
        INSERT INTO kb_sources (
          source_id, path, first_seen_path, sha256, mime, bytes, page_count, sheet_count, status, meta
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
        ON CONFLICT (source_id) DO UPDATE SET
          path = EXCLUDED.path,
          first_seen_path = kb_sources.first_seen_path,
          mime = EXCLUDED.mime,
          bytes = EXCLUDED.bytes,
          page_count = COALESCE(EXCLUDED.page_count, kb_sources.page_count),
          sheet_count = COALESCE(EXCLUDED.sheet_count, kb_sources.sheet_count),
          sha256 = EXCLUDED.sha256,
          status = EXCLUDED.status,
          meta = kb_sources.meta || EXCLUDED.meta
        """,
        (
            source_id,
            path,
            first_seen_path,
            sha256,
            mime,
            bytes_,
            page_count,
            sheet_count,
            status,
            json.dumps(meta),
        ),
    )


def patch_meta_merge(cur: Any, source_id: str, patch: dict[str, Any]) -> None:
    cur.execute(
        "UPDATE kb_sources SET meta = meta || %s::jsonb WHERE source_id = %s",
        (json.dumps(patch), source_id),
    )


def iter_files(root: Path) -> Iterable[Path]:
    if root.is_file():
        yield root
        return
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith(".")]
        for name in filenames:
            if name.startswith("."):
                continue
            yield Path(dirpath) / name


def write_extract_bundle(
    repo_root: Path,
    *,
    source_id: str,
    title: str,
    doc_kind: str,
    chunk_rows: list[dict[str, Any]],
    has_full_text: bool,
) -> Path | None:
    extracts_root = repo_root / "kb_extracts"
    staging_root = extracts_root / ".staging"
    staging_dir = staging_root / source_id
    output_dir = extracts_root / source_id
    if staging_dir.exists():
        shutil.rmtree(staging_dir)
    staging_dir.mkdir(parents=True, exist_ok=True)

    locators = sorted({str(r.get("locator", "")) for r in chunk_rows})
    chunk_ids = [str(r["chunk_id"]) for r in chunk_rows]
    manifest = {
        "manifest_version": 1,
        "source_id": source_id,
        "title": title,
        "doc_kind": doc_kind,
        "has_full_text": has_full_text,
        **({"full_text_file": "full.txt"} if has_full_text else {}),
        "units": [
            {
                "locator": str(r.get("locator", "")),
                "chunk_ids": [r["chunk_id"]],
            }
            for r in chunk_rows
        ],
    }
    (staging_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    if has_full_text:
        full_text = "\n\n".join(str(r.get("content", "")) for r in chunk_rows)
        (staging_dir / "full.txt").write_text(full_text, encoding="utf-8")

    if output_dir.exists():
        shutil.rmtree(output_dir)
    shutil.move(str(staging_dir), str(output_dir))
    return output_dir


def process_one_file(
    *,
    file_path: Path,
    scope: str,
    type_arg: TypeFilter,
    reembed: bool,
    dry_run: bool,
    password: str | None,
    conn: Any | None,
    cur: Any | None,
    client: Any | None,
    embed_batch_size: int,
    embed_dim: int,
    strict: bool,
    repo_root: Path,
    emit_event: bool,
    spend_counter: list[float],
    max_spend: float,
    skip_extract_write: bool,
) -> bool:
    """Returns False if caller should abort (strict failure)."""
    data = file_path.read_bytes()
    if len(data) == 0:
        print(f"skip empty file: {file_path}", flush=True)
        return True

    mime = guess_mime(file_path)
    asset = classify_asset_kind(file_path, mime, type_arg)
    if asset is None:
        print(f"skip (type filter): {file_path}", flush=True)
        return True

    sha = sha256_file(file_path)
    sid = source_id_from_sha256(sha)
    title = title_from_path(file_path)

    first_page_text = ""
    if asset == "pdf":
        pages_peek, st_peek = extract_pdf_pages(data, password)
        if st_peek != "encrypted" and pages_peek:
            first_page_text = pages_peek[0]

    dk = classify_doc_kind(file_path, mime, first_page_text)
    meta_base: dict[str, Any] = {
        "doc_kind": dk.doc_kind,
        "doc_kind_confidence": dk.confidence,
        "doc_kind_matched_by": dk.matched_by,
        "doc_kind_signals": dk.signals,
    }
    if dk.brand:
        meta_base["brand"] = dk.brand
    if dk.product_line:
        meta_base["product_line"] = dk.product_line
    if dk.year is not None:
        meta_base["year"] = dk.year
    if dk.distributor_branded:
        meta_base["distributor_branded"] = True

    if dry_run:
        print(
            json.dumps(
                {
                    "dry_run": True,
                    "path": str(file_path),
                    "source_id": sid,
                    "mime": mime,
                    "asset": asset,
                    "doc_kind": dk.doc_kind,
                },
                indent=2,
            ),
            flush=True,
        )
        return True

    assert conn is not None and cur is not None

    existing = fetch_source_row(cur, sid)
    if existing and existing["status"] == "ready" and existing["sha256"] == sha and not reembed:
        print(f"skip (unchanged, ready): {file_path}", flush=True)
        # Match kb.ingested shape so callers know wiki kb_extracts bundle was not refreshed.
        print(
            json.dumps(
                {"event": "kb.unchanged", "event_version": 1, "source_id": sid},
            ),
            flush=True,
        )
        conn.commit()
        return True

    first_seen_path = str(file_path) if existing is None else str(existing["first_seen_path"])

    upsert_source_row(
        cur,
        source_id=sid,
        path=str(file_path),
        first_seen_path=first_seen_path,
        sha256=sha,
        mime=mime,
        bytes_=len(data),
        page_count=None,
        sheet_count=None,
        status="extracting",
        meta=meta_base,
    )
    conn.commit()

    chunk_rows: list[dict[str, Any]] = []
    failure_reason: str | None = None

    try:
        if asset == "image":
            failure_reason = "needs_vision"
            patch_meta_merge(cur, sid, {"reason": failure_reason})
            upsert_source_row(
                cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=None,
                sheet_count=None,
                status="failed",
                meta={},
            )
            conn.commit()
            return True

        if asset == "structured":
            try:
                units, schema_name = structured_units_from_json(file_path, data.decode("utf-8"))
                meta_base["structured_schema"] = schema_name
                patch_meta_merge(cur, sid, {"structured_schema": schema_name})
                for i, u in enumerate(units):
                    chunk_rows.append(
                        {
                            "chunk_id": u["chunk_id"],
                            "source_id": sid,
                            "scope": scope,
                            "locator": u["locator"],
                            "chunk_index": i,
                            "content": u["body"],
                            "meta": {"kind": "structured"},
                        }
                    )
            except Exception as e:
                reason = "schema_validation_failed" if "schema" in str(e).lower() else "no_schema"
                failure_reason = getattr(e, "message", str(e)) if reason == "schema_validation_failed" else reason
                patch_meta_merge(cur, sid, {"reason": failure_reason})
                upsert_source_row(
                    cur,
                    source_id=sid,
                    path=str(file_path),
                    first_seen_path=first_seen_path,
                    sha256=sha,
                    mime=mime,
                    bytes_=len(data),
                    page_count=None,
                    sheet_count=None,
                    status="failed",
                    meta={},
                )
                conn.commit()
                if strict:
                    return False
                return True

        elif asset == "pdf":
            pages, st = extract_pdf_pages(data, password)
            if st == "encrypted":
                patch_meta_merge(cur, sid, {"reason": "encrypted"})
                upsert_source_row(
                    cur,
                    source_id=sid,
                    path=str(file_path),
                    first_seen_path=first_seen_path,
                    sha256=sha,
                    mime=mime,
                    bytes_=len(data),
                    page_count=None,
                    sheet_count=None,
                    status="failed",
                    meta={},
                )
                conn.commit()
                if strict:
                    return False
                return True
            if st == "needs_ocr":
                patch_meta_merge(cur, sid, {"reason": "needs_ocr"})
                upsert_source_row(
                    cur,
                    source_id=sid,
                    path=str(file_path),
                    first_seen_path=first_seen_path,
                    sha256=sha,
                    mime=mime,
                    bytes_=len(data),
                    page_count=len(pages),
                    sheet_count=None,
                    status="failed",
                    meta={},
                )
                conn.commit()
                if strict:
                    return False
                return True
            chunk_rows = build_pdf_chunk_rows(source_id=sid, title=title, pages=pages, scope=scope)
            upsert_source_row(
                cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=len(pages),
                sheet_count=None,
                status="extracting",
                meta={},
            )

        elif asset == "excel":
            units, sheet_count, failed_sheets = extract_excel_units(sid, data)
            chunk_index = 0
            for u in units:
                header = f"[source={title}][locator={u['locator']}]\n\n"
                chunk_rows.append(
                    {
                        "chunk_id": u["chunk_id"],
                        "source_id": sid,
                        "scope": scope,
                        "locator": u["locator"],
                        "chunk_index": chunk_index,
                        "content": header + u["body"],
                        "meta": {**u["meta"], "failed_sheets": failed_sheets},
                    }
                )
                chunk_index += 1
            meta_patch = {"failed_sheets": failed_sheets} if failed_sheets else {}
            if meta_patch:
                patch_meta_merge(cur, sid, meta_patch)
            upsert_source_row(
                cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=None,
                sheet_count=sheet_count,
                status="extracting",
                meta={},
            )

        elif asset == "pptx":
            slides = extract_pptx_slides(data)
            pptx_slides_extract = slides
            chunk_rows = build_pptx_chunk_rows(source_id=sid, title=title, slides=slides, scope=scope)
            upsert_source_row(
                cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=None,
                sheet_count=None,
                status="extracting",
                meta={},
            )

        else:
            # text / markdown / csv / misc json
            text = data.decode("utf-8", errors="replace")
            strategy = "section-aware" if dk.doc_kind == "master-spec" else "char-window"
            if strategy == "section-aware":
                sections = [s.strip() for s in re.split(r"\n(?=\s*(?:#{1,6}\s+|\d+(?:\.\d+)*\s+[A-Z]))", text) if s.strip()]
                if not sections:
                    sections = [text.strip()]
                parts: list[str] = []
                for sec in sections:
                    parts.extend(chunk_transcript(sec, CHUNK_CHARS, CHUNK_OVERLAP))
            else:
                parts = chunk_transcript(text, CHUNK_CHARS, CHUNK_OVERLAP)
            text_body_extract = text
            for i, part in enumerate(parts):
                header = f"[source={title}]\n\n"
                chunk_rows.append(
                    {
                        "chunk_id": f"{sid}_{i:05d}",
                        "source_id": sid,
                        "scope": scope,
                        "locator": "",
                        "chunk_index": i,
                        "content": header + part,
                        "meta": {"kind": "text"},
                    }
                )
            upsert_source_row(
                cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=None,
                sheet_count=None,
                status="extracting",
                meta={},
            )

        if not chunk_rows:
            patch_meta_merge(cur, sid, {"reason": "empty_extract"})
            upsert_source_row(
                cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=None,
                sheet_count=None,
                status="failed",
                meta={},
            )
            conn.commit()
            if strict:
                return False
            return True

        if reembed:
            delete_chunks_for_source(cur, sid)
            conn.commit()

        upsert_source_row(
            cur,
            source_id=sid,
            path=str(file_path),
            first_seen_path=first_seen_path,
            sha256=sha,
            mime=mime,
            bytes_=len(data),
            page_count=None,
            sheet_count=None,
            status="embedding",
            meta={},
        )
        conn.commit()

        texts = [str(r["content"]) for r in chunk_rows]
        est_tokens = sum(max(1, len(t) // 4) for t in texts)
        proj_cost = (est_tokens / 1000.0) * EMBED_PRICE_USD_PER_1K
        print(f"embed projection ~${proj_cost:.4f} for {file_path.name} ({len(texts)} chunks)", flush=True)

        assert client is not None
        vectors, _tok = embed_batches(client, EMBED_MODEL, texts, embed_batch_size)
        add_cost = (_tok / 1000.0) * EMBED_PRICE_USD_PER_1K
        spend_counter[0] += add_cost
        if spend_counter[0] > max_spend:
            raise SystemExit(f"Exceeded --max-embed-spend cap (${max_spend}); cumulative ~${spend_counter[0]:.4f}")

        sql = """
            INSERT INTO kb_rag_chunks (chunk_id, source_id, scope, locator, chunk_index, content, embedding, status, meta)
            VALUES (%s, %s, %s, %s, %s, %s, %s, 'ready', %s::jsonb)
            ON CONFLICT (chunk_id) DO UPDATE SET
              source_id = EXCLUDED.source_id,
              scope = EXCLUDED.scope,
              locator = EXCLUDED.locator,
              chunk_index = EXCLUDED.chunk_index,
              content = EXCLUDED.content,
              embedding = EXCLUDED.embedding,
              status = 'ready',
              meta = EXCLUDED.meta
        """
        for row, vec in zip(chunk_rows, vectors):
            if len(vec) != embed_dim:
                raise SystemExit(f"embedding dim {len(vec)} != expected {embed_dim}")
            cur.execute(
                sql,
                (
                    row["chunk_id"],
                    row["source_id"],
                    row["scope"],
                    row["locator"],
                    row["chunk_index"],
                    row["content"],
                    vec,
                    json.dumps(row.get("meta") or {}),
                ),
            )

        upsert_source_row(
            cur,
            source_id=sid,
            path=str(file_path),
            first_seen_path=first_seen_path,
            sha256=sha,
            mime=mime,
            bytes_=len(data),
            page_count=None,
            sheet_count=None,
            status="ready",
            meta={},
        )
        conn.commit()

        has_full = dk.doc_kind in FULL_TEXT_DOC_KINDS
        extract_path: Path | None = None
        if not skip_extract_write:
            mf_extra: dict[str, Any] = {}
            if dk.brand:
                mf_extra["brand"] = dk.brand
            if dk.product_line:
                mf_extra["product_line"] = dk.product_line
            if dk.year is not None:
                mf_extra["year"] = dk.year
            if dk.distributor_branded:
                mf_extra["distributor_branded"] = True

            unit_files: dict[str, str] = {}
            units_manifest: list[dict[str, Any]] = []
            full_text_file: str | None = None
            full_text_body: str | None = None
            structured_schema_nm = meta_base.get("structured_schema")

            if asset == "pdf" and pdf_pages_extract is not None:
                for i, pt in enumerate(pdf_pages_extract, start=1):
                    unit_files[f"unit-page-{i:03d}.txt"] = pt
                units_manifest = pdf_manifest_units(pdf_pages_extract, chunk_rows)
                if has_full:
                    full_text_file = "full.txt"
                    full_text_body = "\n\n".join(pdf_pages_extract)

            elif asset == "pptx" and pptx_slides_extract is not None:
                for n, txt in pptx_slides_extract:
                    unit_files[f"unit-slide-{n:03d}.txt"] = txt
                units_manifest = pptx_manifest_units(len(pptx_slides_extract), chunk_rows)
                if has_full:
                    full_text_file = "full.txt"
                    full_text_body = "\n\n".join(t for _, t in pptx_slides_extract)

            elif asset == "excel" and excel_sheet_order is not None and excel_sheet_artifacts is not None:
                sheet_fn_map = {sn: excel_sheet_artifacts[sn][0] for sn in excel_sheet_order if sn in excel_sheet_artifacts}
                for sn in excel_sheet_order:
                    if sn not in excel_sheet_artifacts:
                        continue
                    fn, csv_c = excel_sheet_artifacts[sn]
                    unit_files[fn] = csv_c
                units_manifest = excel_manifest_units(excel_sheet_order, chunk_rows, sheet_fn_map)

            elif asset == "structured":
                for i, r in enumerate(chunk_rows):
                    unit_files[f"unit-record-{i:04d}.txt"] = str(r["content"])
                units_manifest = structured_manifest_units(chunk_rows)

            else:
                if text_body_extract is not None:
                    unit_files["unit-full.txt"] = text_body_extract
                units_manifest = text_manifest_units(chunk_rows)
                if has_full and text_body_extract is not None:
                    full_text_file = "full.txt"
                    full_text_body = text_body_extract

            extract_path = publish_kb_extract_bundle(
                repo_root,
                source_id=sid,
                title=title,
                doc_kind=dk.doc_kind,
                manifest_version=1,
                units=units_manifest,
                has_full_text=bool(has_full and full_text_body),
                full_text_file=full_text_file if (has_full and full_text_body) else None,
                full_text_content=full_text_body if (has_full and full_text_body) else None,
                unit_files=unit_files,
                manifest_fields=mf_extra or None,
                structured_schema=structured_schema_nm if asset == "structured" else None,
            )

        if emit_event:
            evt = {
                "event": "kb.ingested",
                "event_version": 1,
                "source_id": sid,
                "title": title,
                "chunk_ids": [r["chunk_id"] for r in chunk_rows],
                "locators": sorted({str(r.get("locator") or "") for r in chunk_rows}),
                "extracted_text_path": str(extract_path / "full.txt") if extract_path and has_full else "",
            }
            print(json.dumps(evt), flush=True)

    except Exception as e:
        conn.rollback()
        with conn.cursor() as err_cur:
            patch_meta_merge(err_cur, sid, {"reason": str(e)})
            upsert_source_row(
                err_cur,
                source_id=sid,
                path=str(file_path),
                first_seen_path=first_seen_path,
                sha256=sha,
                mime=mime,
                bytes_=len(data),
                page_count=None,
                sheet_count=None,
                status="failed",
                meta={},
            )
        conn.commit()
        print(f"FAILED {file_path}: {e}", flush=True)
        if strict:
            return False
    return True


def parse_args(argv: list[str]) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Ingest raw KB files into Postgres pgvector tables.")
    p.add_argument("path", type=Path, help="File or directory to ingest")
    p.add_argument("--type", dest="file_type", choices=("auto", "pdf", "excel", "text", "image"), default="auto")
    p.add_argument("--scope", default="global", help="Chunk scope (default global)")
    p.add_argument("--reembed", action="store_true", help="Force re-embed and upsert even if sha unchanged")
    p.add_argument("--dry-run", action="store_true", help="Show planned work only; no DB or API calls")
    p.add_argument(
        "--max-embed-spend",
        type=float,
        default=DEFAULT_MAX_EMBED_SPEND_USD,
        help=f"Max cumulative embedding spend in USD (default {DEFAULT_MAX_EMBED_SPEND_USD})",
    )
    p.add_argument(
        "--on-success",
        choices=("emit-event",),
        default=None,
        help='If set to emit-event, print kb.ingested JSON lines after each source succeeds',
    )
    p.add_argument("--batch", type=int, default=64, help="Embedding batch size")
    p.add_argument("--password", default=None, help="Password for encrypted PDFs / workbooks")
    p.add_argument("--strict", action="store_true", help="Abort on first failed source")
    p.add_argument(
        "--skip-extract-files",
        action="store_true",
        help="Skip writing kb_extracts/<source_id>/ artifacts",
    )
    return p.parse_args(argv)


def ensure_kb_schema(conn: Any, embed_dim: int) -> None:
    """Apply schema.sql with the configured embedding dimension."""
    sql = SCHEMA_PATH.read_text(encoding="utf-8")
    sql = sql.replace("vector(1536)", f"vector({embed_dim})")
    with conn.cursor() as cur:
        cur.execute(sql)
    conn.commit()


def main(argv: list[str] | None = None) -> None:
    argv = argv if argv is not None else sys.argv[1:]
    args = parse_args(argv)

    load_dotenv(REPO_ROOT / ".env")
    load_dotenv(CALLS_DIR / "sauron" / ".env")

    root = args.path.resolve()
    if not root.exists():
        raise SystemExit(f"Path not found: {root}")

    embed_dim = int(os.environ.get("PGVECTOR_DIMENSION", str(EMBED_DIM)))

    if args.dry_run:
        for f in sorted(iter_files(root)):
            process_one_file(
                file_path=f,
                scope=args.scope,
                type_arg=args.file_type,
                reembed=args.reembed,
                dry_run=True,
                password=args.password,
                conn=None,
                cur=None,
                client=None,
                embed_batch_size=args.batch,
                embed_dim=embed_dim,
                strict=args.strict,
                repo_root=REPO_ROOT,
                emit_event=False,
                spend_counter=[0.0],
                max_spend=args.max_embed_spend,
                skip_extract_write=True,
            )
        return

    api_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise SystemExit("OPENAI_API_KEY is required (unless using --dry-run).")

    try:
        from openai import OpenAI
        import psycopg
        from pgvector.psycopg import register_vector
    except ImportError as e:
        raise SystemExit(f"Missing dependency: {e}. pip install -r requirements-kb-ingest.txt")

    conninfo = build_conninfo()
    client = OpenAI(api_key=api_key)

    spend_counter = [0.0]
    with psycopg.connect(conninfo, connect_timeout=30) as conn:
        ensure_kb_schema(conn, embed_dim)
        register_vector(conn)
        for f in sorted(iter_files(root)):
            with conn.cursor() as cur:
                ok = process_one_file(
                    file_path=f,
                    scope=args.scope,
                    type_arg=args.file_type,
                    reembed=args.reembed,
                    dry_run=False,
                    password=args.password,
                    conn=conn,
                    cur=cur,
                    client=client,
                    embed_batch_size=args.batch,
                    embed_dim=embed_dim,
                    strict=args.strict,
                    repo_root=REPO_ROOT,
                    emit_event=args.on_success == "emit-event",
                    spend_counter=spend_counter,
                    max_spend=args.max_embed_spend,
                    skip_extract_write=args.skip_extract_files,
                )
            if not ok:
                raise SystemExit(1)


if __name__ == "__main__":
    main()
